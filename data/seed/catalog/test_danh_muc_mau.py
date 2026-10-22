"""Kiểm thử bộ sinh và bộ nạp danh mục mẫu (DL-06). Chạy:
python3 -I -m unittest discover -s data/seed/catalog -p 'test_*.py'
"""
import io
import json
import os
import sys
import tempfile
import threading
import unittest
from contextlib import redirect_stderr, redirect_stdout
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import nap_danh_muc  # noqa: E402
import sinh_danh_muc_mau as s  # noqa: E402


class SinhDanhMucTest(unittest.TestCase):
    def test_nguon_hop_le_du_chin_khoi_va_hai_nhom(self):
        dong = s.doc_nguon()
        self.assertEqual(s.kiem(dong), [])
        self.assertEqual({d["khoi"] for d in dong}, set(s.KHOI))

    def test_bat_tu_mo_ho_khong_phan_biet_dau(self):
        self.assertEqual(s.tu_mo_ho("Tang cuong kiem tra"), "tăng cường")
        self.assertIsNone(s.tu_mo_ho("Kiểm tra trật tự xây dựng"))

    def test_bat_loi_dong(self):
        d = dict(s.doc_nguon()[0])
        d.update(ma="TC-K2-01", muc_kho="6", diem_chuan="101", nhom="Khác", ten="Đẩy mạnh tuyên truyền")
        loi = " | ".join(s.kiem_dong(d))
        for k in ("tiền tố khối", "mức khó", "điểm chuẩn", "nhóm", "mơ hồ"):
            self.assertIn(k, loi)

    def test_tep_excel_co_tieu_de_may_chu_nhan_va_danh_dau_mau(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "k3.xlsx"
            n = s.ghi_xlsx(s.doc_nguon(), "K3", p)
            ws = load_workbook(p).active
            hang = list(ws.iter_rows(values_only=True))
        self.assertEqual(len(hang) - 1, n)
        self.assertEqual(hang[0][1], "Tên công việc")
        self.assertTrue(all(r[-1] == s.DANH_DAU_MAU for r in hang[1:]))
        self.assertIsInstance(hang[1][5], int)

    def test_tep_trong_repo_khop_nguon(self):
        from openpyxl import load_workbook

        for k in s.KHOI:
            ws = load_workbook(HERE / f"danh-muc-mau-{k}.xlsx").active
            ma = [r[0] for r in ws.iter_rows(min_row=2, values_only=True)]
            self.assertEqual(ma, [d["ma"] for d in s.doc_nguon() if d["khoi"] == k], k)


class NapDanhMucTest(unittest.TestCase):
    def setUp(self):
        self.goi = []
        test = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _tra(self, st, body, headers=None):
                raw = json.dumps(body).encode()
                self.send_response(st)
                for k, v in (headers or {}).items():
                    self.send_header(k, v)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)

            def do_GET(self):
                test.goi.append(("GET", self.path, dict(self.headers)))
                if self.path.startswith("/api/v1/catalog?"):
                    self._tra(200, {"items": [{"id": "dm-1", "status": "AP_DUNG"}, {"id": "dm-2", "status": "NHAP"}], "nextCursor": None})
                else:
                    self._tra(200, {"id": "j1", "status": "XONG", "totalRows": 6, "validRows": 6, "errorGroups": []})

            def do_POST(self):
                n = int(self.headers["Content-Length"])
                body = self.rfile.read(n)
                test.goi.append(("POST", self.path, dict(self.headers), body))
                self._tra(202, {"id": "j1", "status": "DANG_XU_LY", "totalRows": 0, "validRows": 0, "errorGroups": []},
                          {"Location": "/api/v1/catalog/imports/j1"})

        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.env = {"MISSION_API": f"http://127.0.0.1:{self.srv.server_port}/api/v1", "MISSION_TOKEN": "t"}

    def tearDown(self):
        self.srv.shutdown()

    def test_nap_chon_ban_nhap_gui_multipart_va_cho_tac_vu(self):
        old = dict(os.environ)
        os.environ.update(self.env)
        nap_danh_muc.time.sleep = lambda _s: None
        try:
            with redirect_stdout(io.StringIO()) as out, redirect_stderr(io.StringIO()):
                rc = nap_danh_muc.main(["K1"])
        finally:
            os.environ.clear()
            os.environ.update(old)
        self.assertEqual(rc, 0, out.getvalue())
        post = [g for g in self.goi if g[0] == "POST"][0]
        self.assertEqual(post[1], "/api/v1/catalog/dm-2/imports")
        self.assertIn("Idempotency-Key", post[2])
        self.assertTrue(post[2]["Content-Type"].startswith("multipart/form-data; boundary="))
        self.assertIn(b'filename="danh-muc-mau-K1.xlsx"', post[3])
        self.assertEqual(self.goi[-1][1], "/api/v1/catalog/imports/j1")
        self.assertIn("Hợp lệ 6/6", out.getvalue())

    def test_dia_chi_tac_vu(self):
        api = "http://h/api/v1"
        self.assertEqual(nap_danh_muc.dia_chi_tac_vu(api, "/api/v1/catalog/imports/x"), "http://h/api/v1/catalog/imports/x")
        self.assertEqual(nap_danh_muc.dia_chi_tac_vu(api, "/catalog/imports/x"), "http://h/api/v1/catalog/imports/x")


if __name__ == "__main__":
    unittest.main()
