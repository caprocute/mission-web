"""Sinh bộ ca mẫu của gói twd/diem-cong-viec từ Phụ lục 01 và Phụ lục 01G của bộ tài liệu bàn giao UEB (DL-16).

Cách chạy (từ gốc repo, hai tệp xlsx lấy từ bộ tài liệu đã giải nén ở máy người chạy, không đưa vào repo):
    python -I data/golden/twd/sinh_ca_twd.py <Phu_luc_01.xlsx> <Phu_luc_01G.xlsx>

Kết quả: data/policy-pack/twd/diem-cong-viec/tests/bo-ca-mau-twd.yaml. Kỳ vọng của từng ca lấy từ giá trị đã
tính sẵn trong bảng tính (cột kết quả), không tính lại bằng công thức của bộ chạy ca, nên bộ chạy ca và gói
được đối chiếu độc lập với Phụ lục 01. Tập lệnh tự kiểm phạm vi phụ trách dựng lại có khớp cột điểm từng
người của bảng tính không; lệch thì dừng. Ca theo Đặc tả 11.12 điểm h (G1 – G13), mục 32.4 và Phụ lục C
lấy kết quả mong đợi ghi trong Đặc tả. Cần Python 3.12, PyYAML, openpyxl.
"""
from __future__ import annotations

import collections
import datetime
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

GOC = Path(__file__).resolve().parents[3]
DICH = GOC / "data/policy-pack/twd/diem-cong-viec/tests/bo-ca-mau-twd.yaml"
MUC = {4: "HTXSNV", 3: "HTTNV", 2: "HTNV", 1: "KHTNV"}
TCC = {"Thực hiện tốt": "THUC_HIEN_TOT", "Đảm bảo": "DAM_BAO", "Cơ bản đảm bảo": "CO_BAN_DAM_BAO", "Không đảm bảo": "KHONG_DAM_BAO"}
TRANG_THAI = {1: "HOAN_THANH", 0: "CHUA_HOAN_THANH", "KHT": "KHONG_DUOC_CHAP_NHAN", "CK": "CHUYEN_KY"}
LY_DO = [("không được giao việc trọng tâm", "co_viec_trong_tam"), ("chưa hoàn thành 100% trọng tâm", "hoan_thanh_100_trong_tam"),
         ("nhiệm vụ, công việc vượt dưới", "dat_ty_le_vuot"), ("chưa có giải pháp đột phá", "giai_phap_dot_pha"),
         ("có nhiệm vụ trọng tâm không hoàn thành do chủ quan", "truot_trong_tam_chu_quan")]
NGUON_01, NGUON_01G, NGUON_DT = "Phụ lục 01 (bản 14/10/2026)", "Phụ lục 01G", "Đặc tả UEB"


def doc(tep: str) -> dict:
    wb = openpyxl.load_workbook(tep, data_only=True, read_only=True)
    kq = {}
    for ws in wb.worksheets:
        dong = {}
        for i, r in enumerate(ws.iter_rows(values_only=True), 1):
            d = {get_column_letter(j + 1): (v.isoformat()[:10] if isinstance(v, datetime.datetime) else v)
                 for j, v in enumerate(r) if v is not None}
            if d:
                dong[i] = d
        kq[ws.title] = dong
    return kq


def so(x, chu_so: int = 4):
    """Làm tròn nửa lên từ giá trị bảng tính; trả int khi là số nguyên để YAML gọn."""
    d = Decimal(repr(float(x or 0))).quantize(Decimal(1).scaleb(-chu_so), rounding=ROUND_HALF_UP)
    return int(d) if d == d.to_integral() else float(d)


def r2r1(x):
    return so(so(x, 2), 1)


def b(x) -> bool:
    return x == 1


class Du:
    def __init__(self, pl01: dict, pl01g: dict):
        self.pl01, self.pl01g = pl01, pl01g
        t = pl01["Cây toàn trình"]
        self.rows = [r for i, r in sorted(t.items()) if i > 4 and ".Q4." in str(r.get("A", ""))]
        self.by = {r["A"]: r for r in self.rows}
        self.kids = collections.defaultdict(list)
        for r in self.rows:
            if r.get("BX"):
                self.kids[r["BX"]].append(r["A"])
        self.txr = [r for i, r in sorted(pl01["Cây thường xuyên"].items()) if i > 4 and r.get("B") == "Công việc"]
        self.nguoi = [r for i, r in sorted(pl01["Điểm quý IV"].items()) if i >= 5 and isinstance(r.get("A"), int)]
        self.cv = [r for r in self.rows if r["G"] == "Công việc (CV)"]
        self.la = [r for r in self.rows if r.get("CB") == 1]
        self.khac_phuc = {r["CJ"] for r in self.rows if r.get("CJ")}

    def truc(self, r: dict) -> str:
        while not r.get("U"):
            r = self.by[r["BX"]]
        return str(r["U"]).split()[0]

    def cvs(self, ma: str) -> list[dict]:
        out = []
        for k in self.kids[ma]:
            if self.by[k]["G"] == "Công việc (CV)":
                out.append(self.by[k])
            elif self.by[k]["G"] != "Công việc con (CVc)":
                out += self.cvs(k)
        return out

    def tong_ql(self, r: dict) -> dict:
        L = [x for x in self.cvs(r["A"]) if x.get("CL") == 1]
        return {"tong_giao": so(sum(x.get("BH") or 0 for x in L), 6), "tong_so_luong": so(sum(x.get("BI") or 0 for x in L), 6),
                "tong_chat_luong": so(sum(x.get("BJ") or 0 for x in L), 6), "tong_tien_do": so(sum(x.get("BK") or 0 for x in L), 6)}

    def ql(self, r: dict) -> list:
        """Điểm quản lý của mục theo bảng tính (cột BT – BW) – là phần đóng góp vào điểm cá nhân."""
        return [r.get(c) or 0 for c in ("BT", "BU", "BV", "BW")]

    def anc(self, ten: str, cot: str, cap: str) -> list[str]:
        out = []
        for r in self.cv:
            if r.get(cot) == ten:
                a = r
                while a["G"] != cap:
                    a = self.by[a["BX"]]
                if a["A"] not in out:
                    out.append(a["A"])
        return out


def muc(ma, nhom, giao_sl_cl_td, truc=None, truc_tiep=False) -> dict:
    g, s, q, d = (so(x, 6) for x in giao_sl_cl_td)
    m = {"ma": ma, "nhom": nhom, "giao": g, "so_luong": s, "chat_luong": q, "tien_do": d}
    if truc:
        m["truc"] = truc
    if truc_tiep:
        m["truc_tiep"] = True
    return m


def cap_cua(chuc_danh: str) -> str:
    for tien_to, cap in (("Cấp trưởng đơn vị", "CAP_TRUONG_DON_VI"), ("Cấp phó đơn vị", "CAP_PHO_DON_VI"),
                         ("Cấp trưởng vụ", "CAP_TRUONG_VU"), ("Cấp phó vụ", "CAP_PHO_VU")):
        if chuc_danh.startswith(tien_to):
            return cap
    return "KHONG_GIU_CHUC_VU"


def pham_vi(du: Du, p: dict) -> tuple[list[dict], dict]:
    """Phạm vi phụ trách của một người theo Đặc tả 11.4, 11.5 (dựng lại cách Phụ lục 01 gom điểm) và phần
    trực tiếp làm của lãnh đạo; trả (các mục, các mục theo vị trí kiêm nhiệm)."""
    ten, cap = p["B"], cap_cua(p["C"])
    out, kiem_nhiem = [], []
    wtx = lambda r: [r.get(c) or 0 for c in ("Y", "Z", "AA", "AB")]
    wcv = lambda r: [r.get(c) or 0 for c in ("BH", "BI", "BJ", "BK")]
    if cap == "CAP_TRUONG_DON_VI":
        out += [muc(r["A"], "THEN_CHOT", wcv(r), du.truc(r)) for r in du.cv]
        out += [muc(r["A"], "THUONG_XUYEN", wtx(r)) for r in du.txr if r["D"] != "DU"]
    elif cap == "CAP_PHO_DON_VI":
        out += [muc(c, "THEN_CHOT", du.ql(du.by[c]), du.truc(du.by[c])) for c in du.anc(ten, "CF", "KR quý")]
        out += [muc(r["A"], "THUONG_XUYEN", wtx(r)) for r in du.txr if r.get("H") == ten]
        kiem_nhiem = [muc(r["A"], "THUONG_XUYEN", wtx(r)) for r in du.txr if r["D"] == "DU" and str(r.get("G")).startswith(ten + " (")]
    elif cap == "CAP_TRUONG_VU":
        out += [muc(c, "THEN_CHOT", du.ql(du.by[c]), du.truc(du.by[c])) for c in du.anc(ten, "CE", "Nhiệm vụ (NV)")]
        out += [muc(r["A"], "THEN_CHOT", wcv(r), du.truc(r)) for r in du.cv if r.get("CG") and r.get("CG") == p["D"]]
        out += [muc(r["A"], "THUONG_XUYEN", wtx(r)) for r in du.txr if r.get("G") == ten]
    elif cap == "CAP_PHO_VU":
        out += [muc(c, "THEN_CHOT", du.ql(du.by[c]), du.truc(du.by[c])) for c in du.anc(ten, "CD", "KPI")]
        out += [muc(r["A"], "THEN_CHOT", wcv(r), du.truc(r)) for r in du.cv if r.get("CG") and r.get("AT") == ten]
        out += [muc(r["A"], "THUONG_XUYEN", wtx(r)) for r in du.txr if r.get("F") == ten]
    else:
        out += [muc(r["A"], "THEN_CHOT", [r.get("BH") or 0] + [r.get(c) or 0 for c in ("BL", "BM", "BN")], du.truc(r))
                for r in du.la if r.get("AS") == ten]
        out += [muc(r["A"], "THUONG_XUYEN", wtx(r)) for r in du.txr if r.get("E") == ten]
    if cap != "KHONG_GIU_CHUC_VU":         # phần lãnh đạo trực tiếp làm (Đặc tả 11.5), chỉ cộng khi tham số bật cho cấp
        out += [muc(r["A"], "THEN_CHOT", [r.get("BH") or 0] + [r.get(c) or 0 for c in ("BL", "BM", "BN")], du.truc(r), True)
                for r in du.la if r.get("AS") == ten]
        out += [muc(r["A"], "THUONG_XUYEN", wtx(r), truc_tiep=True) for r in du.txr if r.get("E") == ten]
    return out + kiem_nhiem, {"ban": out, "kiem_nhiem": kiem_nhiem}


def tong(ds: list[dict], nhom: str, truc_tiep: bool = False) -> list[float]:
    L = [m for m in ds if m["nhom"] == nhom and bool(m.get("truc_tiep")) == truc_tiep]
    return [sum(m[k] for m in L) for k in ("giao", "so_luong", "chat_luong", "tien_do")]


def kiem(dieu_kien: bool, thong_bao: str):
    if not dieu_kien:
        sys.exit("Không khớp Phụ lục 01: " + thong_bao)


def ly_do(s) -> list[str]:
    s = str(s or "")
    if not s.startswith("Thiếu điều kiện Xuất sắc"):
        return []
    s = s.split("Không đủ điều kiện mức Tốt")[0]
    return [ten for cum, ten in LY_DO if cum in s]


def sinh(pl01: dict, pl01g: dict) -> dict:
    du = Du(pl01, pl01g)
    ca = []

    def them(id_, loai, mo_ta, legal_ref, source, dau_vao, mong_doi, tham_so=None):
        c = {"id": id_, "loai": loai, "mo_ta": mo_ta, "legal_ref": legal_ref, "source": source}
        if tham_so:
            c["tham_so"] = tham_so
        c["dau_vao"], c["mong_doi"] = dau_vao, mong_doi
        ca.append(c)

    # ------------------------------------------------------------ 1. Điểm công việc (Đặc tả 3.7, 11.3)
    for r in du.rows:
        if not r["G"].startswith("Công việc"):
            continue
        kiem(r.get("AO") == 1 and r.get("AP") == 1.3, f"hệ số của {r['A']}")
        dv = {"diem_chuan": so(r["AN"]), "muc_do_kho": 2, "nhom": "THEN_CHOT", "trang_thai": TRANG_THAI[r["BB"]]}
        if r["BB"] == 1:
            dv["muc_chat_luong"] = r["BD"]
            if r.get("BF"):
                dv["ngay_cham_chu_quan"] = r["BF"]
        if r["G"] == "Công việc (CV)":
            md = {"diem_cong_viec": so(r["AQ"]), "giao": so(r.get("BH")), "so_luong": so(r.get("BI")),
                  "chat_luong": so(r.get("BJ")), "tien_do": so(r.get("BK"))}
            them(f"VIEC-{r['A'].removeprefix('BTCTW.2026.')}", "VIEC", f"Điểm công việc {r['A']}: {str(r.get('H'))[:120]}",
                 "Đặc tả UEB 3.7, 11.3, 11.4", NGUON_01, dv, md)
        else:
            cha = du.by[r["BX"]]
            dv["cha_hoan_thanh"] = cha.get("BB") == 1
            md = {"diem_cong_viec": so(r["AQ"]), "giao": so(r.get("BH")), "so_luong": so(r.get("BL")),
                  "chat_luong": so(r.get("BM")), "tien_do": so(r.get("BN"))}
            them(f"VIEC-{r['A'].removeprefix('BTCTW.2026.')}", "VIEC",
                 f"Điểm cá nhân của công việc con {r['A']} (công việc cha {'đã' if dv['cha_hoan_thanh'] else 'chưa'} hoàn thành)",
                 "Đặc tả UEB 3.6, 11.3", NGUON_01, dv, md)
    for r in du.txr:
        kiem(r.get("L") == 1 and r.get("M") == 1, f"hệ số của {r['A']}")
        dv = {"diem_chuan": so(r["K"]), "muc_do_kho": 2, "nhom": "THUONG_XUYEN", "trang_thai": TRANG_THAI[r["S"]]}
        if r["S"] == 1:
            dv["muc_chat_luong"] = r["U"]
            if r.get("W"):
                dv["ngay_cham_chu_quan"] = r["W"]
        them(f"VIEC-{r['A'].removeprefix('BTCTW.2026.')}", "VIEC", f"Điểm công việc thường xuyên {r['A']}: {str(r.get('C'))[:120]}",
             "Đặc tả UEB 3.7, 11.3", NGUON_01, dv,
             {"diem_cong_viec": so(r["N"]), "giao": so(r.get("Y")), "so_luong": so(r.get("Z")), "chat_luong": so(r.get("AA")), "tien_do": so(r.get("AB"))})

    # ------------------------------------------------------------ 2. Điểm quản lý (Đặc tả 11.5)
    LOAI = {"KPI": "KPI", "Nhiệm vụ (NV)": "NHIEM_VU", "KR quý": "KR"}
    for r in du.rows:
        if r["G"] not in LOAI:
            continue
        dv = {"loai": LOAI[r["G"]]} | du.tong_ql(r)
        for cot, khoa in (("BQ", "ty_le_dat_chi_tieu"), ("BR", "muc_chat_luong"), ("BS", "ngay_cham_chu_quan")):
            if r.get(cot) not in (None, 1 if cot != "BS" else 0):
                dv[khoa] = r[cot]
        if r["A"] in du.khac_phuc:
            dv["can_khac_phuc"] = True
            kiem((r.get("BR") or 1) <= 0.75, f"mục cần khắc phục {r['A']} có mức chất lượng trên trần")
        them(f"QL-{r['A'].removeprefix('BTCTW.2026.')}", "QUAN_LY", f"Điểm quản lý của {r['G']} {r['A']}",
             "Đặc tả UEB 11.5", NGUON_01, dv,
             {"giao": so(r.get("BT")), "so_luong": so(r.get("BU")), "chat_luong": so(r.get("BV")), "tien_do": so(r.get("BW"))})

    # ------------------------------------------------------------ 3. Điểm quý IV của 120 người
    goi = {r["B"]: r for i, r in pl01g["Điểm quý IV – gói điểm"].items() if i >= 5 and isinstance(r.get("A"), int)}
    phan = {}
    for p in du.nguoi:
        ten = p["B"]
        ds, theo_vt = pham_vi(du, p)
        phan[ten] = (ds, theo_vt, p)
        tt, tx = tong(ds, "THEN_CHOT"), tong(ds, "THUONG_XUYEN")
        kiem(all(abs(a - (p.get(c) or 0)) < 1e-6 for a, c in zip(tt + tx, "FGHIJKLM")), f"phạm vi của {ten}")
        tt_tt, tx_tt = tong(ds, "THEN_CHOT", True), tong(ds, "THUONG_XUYEN", True)
        kiem(abs(tt_tt[0] + tx_tt[0] - (p.get("S") or 0)) < 1e-6, f"phần trực tiếp làm của {ten}")
        dv = {"doi_tuong": "KHONG_GIU_CHUC_VU" if cap_cua(p["C"]) == "KHONG_GIU_CHUC_VU" else "LANH_DAO", "cap": cap_cua(p["C"]), "muc": ds}
        md = {"tt_giao": so(p.get("F")), "tt_so_luong": so(p.get("G")), "tt_chat_luong": so(p.get("H")), "tt_tien_do": so(p.get("I")),
              "tx_giao": so(p.get("J")), "tx_so_luong": so(p.get("K")), "tx_chat_luong": so(p.get("L")), "tx_tien_do": so(p.get("M")),
              "tong_diem": so(p["P"], 2), "diem_xep_loai": r2r1(p["P"])}
        them(f"CN-{p['A']:03d}", "CA_NHAN", f"Điểm quý IV/2026 của {ten} theo tỷ lệ có hệ số", "Đặc tả UEB 11.4, 11.5; 11.8",
             NGUON_01, dv, md)
        if p.get("S"):
            them(f"CN-{p['A']:03d}-TT", "CA_NHAN", f"Điểm quý IV/2026 của {ten} khi bật cộng phần trực tiếp làm cho cấp của người",
                 "Đặc tả UEB 11.5", NGUON_01, dv, {"tong_diem": so(p["U"], 2)}, {"truc_tiep_lam.cap_ap_dung": [dv["cap"]]})
        g = goi[ten]
        them(f"GD-{p['A']:03d}", "CA_NHAN", f"Điểm quý IV/2026 của {ten} theo gói điểm 80 – 20", "Đặc tả UEB 11.12",
             NGUON_01G, dv, {"tt_kpi": None if not g.get("F") else so(g["J"]), "tx_kpi": None if not g.get("H") else so(g["K"]),
                             "tong_diem": so(g["P"], 2)}, {"thang_100.che_do": "GOI_DIEM"})

    # ------------------------------------------------------------ 4. Điểm năm 2026, hai cách dùng kết quả các quý trước
    nam = {r["B"]: r for i, r in pl01["Điểm năm 2026"].items() if i >= 5 and isinstance(r.get("A"), int)}
    xl = {r["B"]: r for i, r in pl01["Xếp loại cá nhân"].items() if 5 <= i <= 124}
    for p in du.nguoi:
        n, x = nam[p["B"]], xl[p["B"]]
        dv = {"quy_tren_phan_mem": [{"giao": so(n["M"], 6), "dat": so(n["N"], 6)}],
              "quy_truoc": [{"giao": so(n[g], 6), "ty_le_dat": so(n[t], 6)} for g, t in (("D", "E"), ("G", "H"), ("J", "K"))]}
        kiem(so(x["D"], 1) == r2r1(n["X"]) and so(x["Z"], 1) == r2r1(n["W"]), f"điểm năm của {p['B']}")
        them(f"NAM-{p['A']:03d}-A", "NAM_DIEM", f"Điểm năm 2026 của {p['B']}, chỉ các quý có trên phần mềm (giá trị thí điểm)",
             "Đặc tả UEB 11.4; 18.5", NGUON_01, dv, {"tong_diem_nam": so(n["W"], 2), "diem_xep_loai": r2r1(n["W"])})
        them(f"NAM-{p['A']:03d}-B", "NAM_DIEM", f"Điểm năm 2026 của {p['B']}, cộng các quý trước (giả định)",
             "Đặc tả UEB 11.4; 13.44", NGUON_01, dv, {"tong_diem_nam": so(n["X"], 2), "diem_xep_loai": r2r1(n["X"])},
             {"nam.cach_dung_quy_truoc": "CONG_QUY_TRUOC"})

    # ------------------------------------------------------------ 5. Chuỗi xếp loại năm 2026 của 120 người và 15 tập thể
    tcc = {r["B"]: r for i, r in pl01["Tiêu chí chung"].items() if 6 <= i <= 125}
    tcc_tt = {r["B"]: r for i, r in pl01["Tiêu chí chung"].items() if i >= 129}
    ma_a, ma_b = {}, {}
    for p in du.nguoi:
        x, t = xl[p["B"]], tcc[p["B"]]
        kiem(TCC[x["E"]] == TCC[t["S"]], f"tiêu chí chung của {p['B']}")
        co_so = {"ty_le_dap_ung_tieu_chi_chung": so(t["R"], 6), "co_viec_trong_tam": b(x["F"]), "hoan_thanh_100_trong_tam": b(x["G"]),
                 "ty_le_vuot": so(x.get("H"), 10), "giai_phap_dot_pha": b(x["I"]), "truot_trong_tam_chu_quan": b(x["J"]),
                 "co_quy_khong_hoan_thanh": b(x["K"])}
        ma_b[p["B"]] = f"XL-{p['A']:03d}-B"
        them(ma_b[p["B"]], "XEP_LOAI_CHUOI", f"Mức gợi ý năm 2026 của {p['B']}, điểm năm cộng các quý trước",
             "Đặc tả UEB 18.1 – 18.6; dự thảo Điều 11, 14, 15", NGUON_01, {"diem_xep_loai": so(x["D"], 1)} | co_so,
             {"muc_tieu_chi_chung": TCC[x["E"]], "tran_theo_diem": MUC[x["L"]], "tran_theo_tieu_chi_chung": MUC[x["M"]],
              "muc_sau_buoc_3": MUC[x["N"]], "muc_sau_dieu_kien_xuat_sac": MUC[x["O"]], "muc_sau_dieu_kien_tot": MUC[x["P"]],
              "muc_goi_y": MUC[x["Q"]], "dieu_kien_thieu_xuat_sac": ly_do(x.get("S"))})
        ma_a[p["B"]] = f"XL-{p['A']:03d}-A"
        them(ma_a[p["B"]], "XEP_LOAI_CHUOI", f"Mức gợi ý năm 2026 của {p['B']}, điểm năm theo giá trị thí điểm",
             "Đặc tả UEB 18.1 – 18.6; dự thảo Điều 11, 14, 15", NGUON_01, {"diem_xep_loai": so(x["Z"], 1)} | co_so,
             {"tran_theo_diem": MUC[x["AA"]], "muc_goi_y": MUC[x["AB"]]})
    ma_tt = []
    for i, r in sorted(pl01["Tập thể"].items()):
        if i < 39 or not isinstance(r.get("A"), int):
            continue
        t = tcc_tt[r["B"]]
        ma_tt.append(f"TT-{r['A']:02d}")
        them(ma_tt[-1], "XEP_LOAI_CHUOI", f"Mức gợi ý năm 2026 của tập thể {r['B']}", "Đặc tả UEB 18.5, 19; dự thảo Điều 11",
             NGUON_01, {"diem_xep_loai": so(r["D"], 1), "ty_le_dap_ung_tieu_chi_chung": so(t["R"], 6), "co_viec_trong_tam": b(r["F"]),
                        "hoan_thanh_100_trong_tam": b(r["G"]), "ty_le_vuot": so(r.get("H"), 10), "giai_phap_dot_pha": b(r["I"]),
                        "truot_trong_tam_chu_quan": b(r["J"])},
             {"muc_tieu_chi_chung": TCC[r["E"]], "tran_theo_diem": MUC[r["K"]], "tran_theo_tieu_chi_chung": MUC[r["L"]],
              "muc_sau_buoc_3": MUC[r["M"]], "muc_sau_dieu_kien_xuat_sac": MUC[r["N"]], "muc_sau_dieu_kien_tot": MUC[r["O"]],
              "muc_goi_y": MUC[r["P"]], "dieu_kien_thieu_xuat_sac": ly_do(r.get("R"))})

    def dem(ds):
        c = collections.Counter(ds)
        return {m: c.get(m, 0) for m in ("HTXSNV", "HTTNV", "HTNV", "KHTNV")}

    them("LO-B", "XEP_LOAI_LO", "Tổng hợp mức gợi ý năm 2026 của 120 người, điểm năm cộng các quý trước", "Đặc tả UEB 32.2",
         NGUON_01, {"ca": list(ma_b.values())}, {"dem_muc": dem(MUC[x["Q"]] for x in xl.values()), "ca_doi_muc": []})
    them("LO-A", "XEP_LOAI_LO", "Tình huống 1 mục 32.4: điểm năm theo giá trị thí điểm, mức gợi ý 3 / 113 / 4 / 0", "Đặc tả UEB 32.4 tình huống 1",
         NGUON_DT, {"ca": list(ma_a.values())}, {"dem_muc": {"HTXSNV": 3, "HTTNV": 113, "HTNV": 4, "KHTNV": 0}, "ca_doi_muc": []})
    them("LO-TT", "XEP_LOAI_LO", "Tổng hợp mức gợi ý của 15 tập thể: 2 / 13 / 0 / 0", "Đặc tả UEB 32.3",
         NGUON_DT, {"ca": ma_tt}, {"dem_muc": {"HTXSNV": 2, "HTTNV": 13, "HTNV": 0, "KHTNV": 0}, "ca_doi_muc": []})
    them("LO-TH2", "XEP_LOAI_LO", "Tình huống 2 mục 32.4: tắt tham số đối tượng không có việc trọng tâm, không mức nào đổi",
         "Đặc tả UEB 32.4 tình huống 2; 18.2", NGUON_DT, {"ca": list(ma_a.values())},
         {"dem_muc": {"HTXSNV": 3, "HTTNV": 113, "HTNV": 4, "KHTNV": 0}, "ca_doi_muc": []},
         {"xep_loai.chan_xuat_sac_khi_khong_co_trong_tam": False})
    nam_nguoi = ["Chuyên viên Vụ V-2", "Phó Vụ trưởng Vụ V-1", "Vụ trưởng Vụ V", "Phó Trưởng Ban A", "Trưởng Ban"]
    for cach, ma, kq in (("B", ma_b, {"HTXSNV": 3, "HTTNV": 110, "HTNV": 7, "KHTNV": 0}), ("A", ma_a, {"HTXSNV": 3, "HTTNV": 108, "HTNV": 9, "KHTNV": 0})):
        them(f"LO-TH3-{cach}", "XEP_LOAI_LO",
             f"Tình huống 3 mục 32.4: công việc 12.Q4.KR2.NV1.K1.C1 chưa xong do chủ quan, năm người trong phạm vi hạ mức ({'cộng các quý trước' if cach == 'B' else 'giá trị thí điểm'})",
             "Đặc tả UEB 32.4 tình huống 3; dự thảo Điều 11 khoản 2.2", NGUON_DT,
             {"ca": list(ma.values()), "thay_doi": {ma[n]: {"truot_trong_tam_chu_quan": True} for n in nam_nguoi}},
             {"dem_muc": kq, "ca_doi_muc": sorted((ma[n] for n in nam_nguoi), key=list(ma.values()).index)})
    xg = {r["B"]: r for i, r in pl01g["Xếp loại – gói điểm"].items() if i >= 5 and isinstance(r.get("A"), int)}
    them("LO-G13", "XEP_LOAI_LO", "Ca G13: mức gợi ý năm 2026 của 120 người khi tính theo gói điểm 80 – 20, không người nào đổi mức",
         "Đặc tả UEB 11.12 điểm h", NGUON_DT,
         {"ca": list(ma_a.values()), "thay_doi": {ma_a[n]: {"diem_xep_loai": so(r["D"], 1)} for n, r in xg.items()}},
         {"dem_muc": {"HTXSNV": 3, "HTTNV": 113, "HTNV": 4, "KHTNV": 0}, "ca_doi_muc": []})
    return {"ca": ca, "phan": phan, "ma_a": ma_a, "ma_b": ma_b}


# ---------------------------------------------------------------------------- Ca theo Đặc tả (kết quả mong đợi ghi trong Đặc tả)
def ca_dac_ta(kq: dict) -> list[dict]:
    phan = kq["phan"]
    ca = []

    def them(id_, loai, mo_ta, legal_ref, dau_vao, mong_doi, tham_so=None, source=NGUON_DT):
        c = {"id": id_, "loai": loai, "mo_ta": mo_ta, "legal_ref": legal_ref, "source": source}
        if tham_so:
            c["tham_so"] = tham_so
        c["dau_vao"], c["mong_doi"] = dau_vao, mong_doi
        ca.append(c)

    def nguoi(ten, muc=None):
        ds, theo_vt, p = phan[ten]
        cap = cap_cua(p["C"])
        return {"doi_tuong": "KHONG_GIU_CHUC_VU" if cap == "KHONG_GIU_CHUC_VU" else "LANH_DAO", "cap": cap, "muc": ds if muc is None else muc}

    GD = {"thang_100.che_do": "GOI_DIEM"}
    LR = "Đặc tả UEB 11.12 điểm h"
    them("G1", "CA_NHAN", "Ca G1 (như ca thử tay CTT11): Chuyên viên Vụ V-5 theo gói 80 – 20", LR + "; Phụ lục C.2 CTT11",
         nguoi("Chuyên viên Vụ V-5"), {"tt_giao": 46.8, "tt_dat": 38.805, "tong_diem": 86.33}, GD)
    them("G1-TL", "CA_NHAN", "Ca G1, đối chiếu: Chuyên viên Vụ V-5 theo tỷ lệ có hệ số", LR, nguoi("Chuyên viên Vụ V-5"), {"tong_diem": 83.62})
    them("G2", "CA_NHAN", "Ca G2: Chuyên viên Vụ Tổ chức - Điều lệ-4, một việc thường xuyên 3 điểm không được chấp nhận", LR + "; điểm g",
         nguoi("Chuyên viên Vụ Tổ chức - Điều lệ-4"), {"tx_giao": 3, "tx_dat": 0, "tong_diem": 80}, GD)
    them("G2-TL", "CA_NHAN", "Ca G2, đối chiếu theo tỷ lệ có hệ số", LR, nguoi("Chuyên viên Vụ Tổ chức - Điều lệ-4"), {"tong_diem": 94.91})
    them("G3", "CA_NHAN", "Ca G3: Phó Vụ trưởng Vụ Tổ chức - Điều lệ-1 (nhóm lãnh đạo) theo gói 80 – 20", LR,
         nguoi("Phó Vụ trưởng Vụ Tổ chức - Điều lệ-1"),
         {"tt_giao": 159.9, "tt_dat": 133.9, "tx_giao": 20, "tx_dat": 20, "tt_kpi": 0.8374, "tx_kpi": 1, "tong_diem": 86.99}, GD)
    them("G3-TL", "CA_NHAN", "Ca G3, đối chiếu theo tỷ lệ có hệ số", LR, nguoi("Phó Vụ trưởng Vụ Tổ chức - Điều lệ-1"), {"tong_diem": 85.55})
    them("G4", "CA_NHAN", "Ca G4: Chuyên viên Vụ Cán bộ-2 chỉ có việc thường xuyên, nhóm thường xuyên có điểm tối đa 100", LR,
         nguoi("Chuyên viên Vụ Cán bộ-2"), {"tx_giao": 22, "tt_kpi": None, "diem_toi_da_tt": 0, "diem_toi_da_tx": 100, "tong_diem": 100}, GD)
    ds_f, vt_f, _ = phan["Phó Trưởng Ban F"]
    them("G5", "CA_NHAN", "Ca G5: Phó Trưởng Ban F, cộng theo nhóm trên mọi vị trí (Ban và Đảng ủy cơ quan)", LR + "; 11.4",
         nguoi("Phó Trưởng Ban F"), {"tt_giao": 96.2, "tx_giao": 55, "tong_diem": 96.43}, GD)
    them("G5-TL", "CA_NHAN", "Ca G5, đối chiếu theo tỷ lệ có hệ số", LR, nguoi("Phó Trưởng Ban F"), {"tong_diem": 95.87})
    them("G5-BAN", "CA_NHAN", "Ca G5: điểm của Phó Trưởng Ban F ở vị trí Phó Trưởng Ban", LR, nguoi("Phó Trưởng Ban F", vt_f["ban"]),
         {"tx_giao": 30, "tong_diem": 95.66}, GD)
    them("G5-DU", "CA_NHAN", "Ca G5: điểm của Phó Trưởng Ban F ở vị trí kiêm nhiệm Bí thư Đảng ủy, chỉ có việc thường xuyên", LR,
         nguoi("Phó Trưởng Ban F", vt_f["kiem_nhiem"]), {"tx_giao": 25, "tong_diem": 98.33}, GD)
    TRUC = GD | {"ty_trong_truc": {"TR1": 0.3, "TR4": 0.7}}
    them("G6", "CA_NHAN", "Ca G6: tỷ trọng trục 1 30%, trục 4 70%; Chuyên viên Vụ Tổ chức - Điều lệ-2 chỉ có việc trọng tâm", LR + "; điểm c",
         nguoi("Chuyên viên Vụ Tổ chức - Điều lệ-2"), {"tt_kpi": 0.825, "diem_toi_da_tt": 100, "tong_diem": 82.5}, TRUC)
    them("G6-KHONG", "CA_NHAN", "Ca G6, đối chiếu: không đặt tỷ trọng theo trục", LR, nguoi("Chuyên viên Vụ Tổ chức - Điều lệ-2"), {"tong_diem": 84}, GD)
    them("G6-V5", "CA_NHAN", "Ca G6: Chuyên viên Vụ V-5 chỉ có việc ở trục 4, tổng điểm không đổi", LR, nguoi("Chuyên viên Vụ V-5"), {"tong_diem": 86.33}, TRUC)
    them("G7", "CA_NHAN", "Ca G7: Vụ trưởng Vụ Tổ chức - Điều lệ, tổng điểm làm tròn từ tổng chưa làm tròn (93,37, không phải 93,38)", LR + "; điểm b",
         nguoi("Vụ trưởng Vụ Tổ chức - Điều lệ"), {"tt_giao": 438.1, "tx_giao": 42, "tt_kpi": 0.9357, "tx_kpi": 0.9258, "tong_diem": 93.37}, GD)
    them("G7-TL", "CA_NHAN", "Ca G7, đối chiếu theo tỷ lệ có hệ số", LR, nguoi("Vụ trưởng Vụ Tổ chức - Điều lệ"), {"tong_diem": 93.48})
    G10 = GD | {"goi_diem.theo_doi_tuong": {"LANH_DAO": {"THEN_CHOT": 80, "THUONG_XUYEN": 20}, "KHONG_GIU_CHUC_VU": {"THEN_CHOT": 70, "THUONG_XUYEN": 30}}}
    them("G10", "CA_NHAN", "Ca G10: gói lãnh đạo 80 – 20, không giữ chức vụ 70 – 30; Chuyên viên Vụ V-5", LR + "; điểm d",
         nguoi("Chuyên viên Vụ V-5"), {"diem_toi_da_tt": 70, "tong_diem": 88.04}, G10)
    them("G10-PVT", "CA_NHAN", "Ca G10: Phó Vụ trưởng Vụ Tổ chức - Điều lệ-1 vẫn theo gói lãnh đạo", LR,
         nguoi("Phó Vụ trưởng Vụ Tổ chức - Điều lệ-1"), {"diem_toi_da_tt": 80, "tong_diem": 86.99}, G10)
    TT_ON = {"truc_tiep_lam.cap_ap_dung": ["CAP_PHO_VU"]}
    them("G11", "CA_NHAN", "Ca G11: như G3, bật cộng phần trực tiếp làm cho cấp Phó Vụ trưởng (một việc thường xuyên 10 điểm)", LR + "; 11.5",
         nguoi("Phó Vụ trưởng Vụ Tổ chức - Điều lệ-1"), {"tx_giao": 30, "tx_dat": 30, "tx_kpi": 1, "tong_diem": 86.99}, GD | TT_ON)
    them("G11-TL", "CA_NHAN", "Ca G11, đối chiếu theo tỷ lệ có hệ số khi cộng phần trực tiếp làm", LR,
         nguoi("Phó Vụ trưởng Vụ Tổ chức - Điều lệ-1"), {"tong_diem": 86.31}, TT_ON)

    # --- Điểm công việc, ca thử tay (Phụ lục C) có kết quả tính được
    V = {"diem_chuan": 10, "muc_do_kho": 2, "nhom": "THEN_CHOT", "trang_thai": "HOAN_THANH", "muc_chat_luong": 1}
    PC = "Đặc tả UEB Phụ lục C"
    them("CTT1", "VIEC", "Ca CTT1: công việc trọng tâm điểm chuẩn 10, mức khó 4 (10 × 1,6 × 1,3)", PC + ".2 CTT1; 3.7", V | {"muc_do_kho": 4},
         {"diem_cong_viec": 20.8, "giao": 20.8})
    them("CTT4", "VIEC", "Ca CTT4: chậm 1 ngày làm việc do chủ quan, tiến độ còn 95% điểm công việc", PC + ".2 CTT4; 11.3",
         V | {"ngay_cham_chu_quan": 1}, {"diem_cong_viec": 13, "tien_do": 12.35})
    them("CTT10", "VIEC", "Ca CTT10: việc ngoài phạm vi phụ trách không vào điểm được giao, điểm đạt của ai", PC + ".2 CTT10; 11.4",
         V | {"ngoai_pham_vi": True}, {"giao": 0, "so_luong": 0, "diem_dat": 0})
    them("CTB12-BT", "VIEC", "Ca CTB12, biến thể: hủy phần công việc con sau hạn, phần đó vẫn trong điểm được giao, điểm đạt bằng 0",
         PC + ".3 CTB12; 7.4, 7.5", V | {"trang_thai": "CHUA_HOAN_THANH"}, {"giao": 13, "diem_dat": 0})
    them("CTB17", "VIEC", "Ca CTB17: đổi giá trị hệ số mức khó 2 trong kỳ (giá trị giả định 1,1), việc đã hoàn thành được tính lại",
         PC + ".3 CTB17; 3.7", V, {"diem_cong_viec": 14.3, "giao": 14.3}, {"he_so.muc_do_kho": {1: 0.8, 2: 1.1, 3: 1.25, 4: 1.6, 5: 2.0}})
    them("CTB19-III", "VIEC", "Ca CTB19 (iii): công việc Không hoàn thành giữ nguyên điểm được giao, điểm đạt bằng 0", PC + ".3 CTB19",
         V | {"trang_thai": "KHONG_DUOC_CHAP_NHAN"}, {"giao": 13, "so_luong": 0, "chat_luong": 0, "tien_do": 0})
    them("CTB29", "VIEC", "Ca CTB29: gia hạn được duyệt, nộp chậm 1 ngày làm việc so với hạn mới, tiến độ trừ 5%", PC + ".3 CTB29",
         V | {"ngay_cham_chu_quan": 1}, {"tien_do": 12.35})
    them("CTB29-I", "VIEC", "Ca CTB29 (i): gia hạn bị từ chối, chậm 4 ngày làm việc, tiến độ trừ 20%", PC + ".3 CTB29",
         V | {"ngay_cham_chu_quan": 4}, {"tien_do": 10.4})
    them("CTB30", "VIEC", "Ca CTB30: đổi loại sản phẩm sang loại có điểm chuẩn 15 sau khi kế hoạch quý được duyệt, điểm công việc tính lại",
         PC + ".3 CTB30; 7.2", V | {"diem_chuan": 15}, {"diem_cong_viec": 19.5, "giao": 19.5})
    them("VIEC-BS", "VIEC", "Bật một ô hệ số bổ sung (hệ số cấp bách giả định 1,2) và mức chất lượng 0,25", "Đặc tả UEB 3.7; Phụ lục B.3 dòng 7",
         V | {"he_so_bo_sung": {"BO_SUNG_1": 1.2}, "muc_chat_luong": 0.25}, {"diem_cong_viec": 15.6, "chat_luong": 3.9},
         {"he_so.bo_sung": [{"ma": "BO_SUNG_1", "bat": True}, {"ma": "BO_SUNG_2", "bat": False}, {"ma": "BO_SUNG_3", "bat": False}]}, source="BA")
    them("VIEC-CK", "VIEC", "Công việc đã chuyển kỳ không vào điểm được giao", "Đặc tả UEB 11.4", V | {"trang_thai": "CHUYEN_KY"}, {"giao": 0}, source="BA")
    Q = {"loai": "KPI", "tong_giao": 50, "tong_so_luong": 50, "tong_chat_luong": 50, "tong_tien_do": 50}
    them("CTB15-II", "QUAN_LY", "Ca CTB15 (ii): chế độ chặt, KPI chưa xác nhận hoàn thành khi khóa sổ, r = 0,9, m = 1, n = 0",
         PC + ".3 CTB15; 11.5", Q | {"ty_le_dat_chi_tieu": 0.9}, {"so_luong": 45, "chat_luong": 50, "tien_do": 50})
    them("QL-KP", "QUAN_LY", "Mục cần khắc phục: mức chất lượng không cao hơn 0,75 dù người xác nhận ghi 1", "Đặc tả UEB 3.5; Phụ lục B.4 dòng 4",
         Q | {"can_khac_phuc": True}, {"muc_chat_luong_ap_dung": 0.75, "chat_luong": 37.5}, source="BA")
    them("QL-DV", "QUAN_LY", "Điểm của đơn vị không áp phần giảm theo tỷ lệ đạt, mức chất lượng, ngày chậm của mục", "Đặc tả UEB 11.5",
         Q | {"loai": "DON_VI", "ty_le_dat_chi_tieu": 0.5, "muc_chat_luong": 0.5, "ngay_cham_chu_quan": 3},
         {"so_luong": 50, "chat_luong": 50, "tien_do": 50}, source="BA")
    them("QL-N", "QUAN_LY", "Nhiệm vụ chậm 2 ngày làm việc do chủ quan, mức chất lượng 0,75", "Đặc tả UEB 11.5",
         Q | {"loai": "NHIEM_VU", "muc_chat_luong": 0.75, "ngay_cham_chu_quan": 2}, {"chat_luong": 37.5, "tien_do": 45}, source="BA")

    # --- Chuỗi xếp loại: tình huống 32.4 và ca phủ từng dòng bảng quyết định
    XS = {"diem_xep_loai": 95, "ty_le_dap_ung_tieu_chi_chung": 1, "co_viec_trong_tam": True, "hoan_thanh_100_trong_tam": True,
          "ty_le_vuot": 0.35, "giai_phap_dot_pha": True, "truot_trong_tam_chu_quan": False}
    LR18 = "Đặc tả UEB 18.5"
    them("TH4", "XEP_LOAI_CHUOI", "Tình huống 4 mục 32.4: Chuyên viên Vụ Cán bộ-1 có một tiêu chí Không đảm bảo (9/10), trần theo tiêu chí chung là Tốt",
         "Đặc tả UEB 32.4 tình huống 4", {"diem_xep_loai": 100, "ty_le_dap_ung_tieu_chi_chung": 0.9, "co_viec_trong_tam": True,
                                          "hoan_thanh_100_trong_tam": True, "ty_le_vuot": 0, "giai_phap_dot_pha": False, "truot_trong_tam_chu_quan": False},
         {"muc_tieu_chi_chung": "DAM_BAO", "tran_theo_tieu_chi_chung": "HTTNV", "muc_goi_y": "HTTNV"})
    them("TH20-VT", "XEP_LOAI_CHUOI", "Tình huống 20 mục 32.4: tập thể Vụ Cán bộ không hoàn thành nhiệm vụ trọng tâm nào, Vụ trưởng xếp mức thấp nhất",
         "Đặc tả UEB 32.4 tình huống 20; 18.4; dự thảo Điều 14 khoản 3.3", XS | {"rang_buoc_tap_the_thap_nhat": True}, {"muc_goi_y": "KHTNV"})
    them("CTT14-PVT", "XEP_LOAI_CHUOI", "Ca CTT14: Phó Vụ trưởng Vụ Cán bộ đủ điều kiện Xuất sắc nhưng tập thể không hoàn thành mục tiêu trọng tâm nào",
         PC + ".2 CTT14; 32.4 tình huống 20", XS | {"rang_buoc_tap_the_ha_xuat_sac": True},
         {"muc_sau_dieu_kien_tot": "HTXSNV", "muc_goi_y": "HTTNV"})
    them("CTT13", "XEP_LOAI_CHUOI", "Ca CTT13: lãnh đạo đủ trần Xuất sắc, có công việc trọng tâm trực tiếp làm không hoàn thành do chủ quan, hạ hai mức",
         PC + ".2 CTT13; " + LR18 + " bước 4", XS | {"truot_trong_tam_chu_quan": True},
         {"muc_sau_buoc_3": "HTXSNV", "muc_sau_dieu_kien_xuat_sac": "HTTNV", "muc_sau_dieu_kien_tot": "HTNV", "muc_goi_y": "HTNV",
          "dieu_kien_thieu_xuat_sac": ["truot_trong_tam_chu_quan"]})
    them("XL-XS", "XEP_LOAI_CHUOI", "Đủ mọi điều kiện Xuất sắc", LR18, XS, {"muc_goi_y": "HTXSNV", "dieu_kien_thieu_xuat_sac": []}, source="BA")
    them("XL-DR09-TAT", "XEP_LOAI_CHUOI", "Không có việc trọng tâm, tắt tham số chặn, đủ các điều kiện còn lại thì được Xuất sắc",
         "Đặc tả UEB 18.2; 32.4 tình huống 2", XS | {"co_viec_trong_tam": False}, {"muc_goi_y": "HTXSNV"},
         {"xep_loai.chan_xuat_sac_khi_khong_co_trong_tam": False}, source="BA")
    them("XL-DR09-BAT", "XEP_LOAI_CHUOI", "Không có việc trọng tâm, tham số chặn đang bật thì không được Xuất sắc", "Đặc tả UEB 18.2",
         XS | {"co_viec_trong_tam": False}, {"muc_goi_y": "HTTNV", "dieu_kien_thieu_xuat_sac": ["co_viec_trong_tam"]}, source="BA")
    them("XL-VUOT-BIEN", "XEP_LOAI_CHUOI", "Tỷ lệ việc vượt đúng 30% đạt ngưỡng tối thiểu; 29,99% thì không", "Đặc tả UEB 18.2",
         XS | {"ty_le_vuot": 0.3}, {"dat_ty_le_vuot": True, "muc_goi_y": "HTXSNV"}, source="BA")
    them("XL-VUOT-DUOI", "XEP_LOAI_CHUOI", "Tỷ lệ việc vượt 29,99%, dưới ngưỡng", "Đặc tả UEB 18.2",
         XS | {"ty_le_vuot": 0.2999}, {"dat_ty_le_vuot": False, "muc_goi_y": "HTTNV", "dieu_kien_thieu_xuat_sac": ["dat_ty_le_vuot"]}, source="BA")
    them("XL-LOAI-TRU", "XEP_LOAI_CHUOI", "Có điều kiện loại trừ thì xếp mức Không hoàn thành nhiệm vụ", LR18 + " bước 5",
         XS | {"loai_tru": True}, {"muc_goi_y": "KHTNV"}, source="BA")
    them("XL-DR08", "XEP_LOAI_CHUOI", "Có quý Không hoàn thành nhiệm vụ thì không xếp Xuất sắc cả năm", LR18 + " bước 6",
         XS | {"co_quy_khong_hoan_thanh": True}, {"muc_sau_dieu_kien_tot": "HTXSNV", "muc_goi_y": "HTTNV"}, source="BA")
    for id_, mo_ta, d, t, b3, cuoi in (
            ("XL-B3-1", "Dưới 60 điểm, tiêu chí chung Thực hiện tốt", 55, 1, "KHTNV", "KHTNV"),
            ("XL-B3-2", "95 điểm, tiêu chí chung Không đảm bảo (dưới 60%)", 95, 0.5, "KHTNV", "KHTNV"),
            ("XL-B3-3", "70 điểm, tiêu chí chung Thực hiện tốt", 70, 1, "HTNV", "HTNV"),
            ("XL-B3-4", "95 điểm, tiêu chí chung Cơ bản đảm bảo (80%)", 95, 0.8, "HTNV", "HTNV"),
            ("XL-B3-5", "85 điểm, tiêu chí chung Thực hiện tốt", 85, 1, "HTTNV", "HTTNV"),
            ("XL-B3-6", "95 điểm, tiêu chí chung Đảm bảo (90%)", 95, 0.9, "HTTNV", "HTTNV"),
            ("XL-B3-7", "85 điểm, có việc trọng tâm không hoàn thành do chủ quan", 85, 1, "HTTNV", "HTNV")):
        them(id_, "XEP_LOAI_CHUOI", f"Bước 3, lấy mức thấp hơn của hai trần: {mo_ta}", LR18 + " bước 1 – 4",
             XS | {"diem_xep_loai": d, "ty_le_dap_ung_tieu_chi_chung": t, "truot_trong_tam_chu_quan": id_ == "XL-B3-7"},
             {"muc_sau_buoc_3": b3, "muc_goi_y": cuoi, "dieu_kien_thieu_xuat_sac": []}, source="BA")
    return ca


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__)
        return 2
    pl01, pl01g = doc(argv[0]), doc(argv[1])
    kq = sinh(pl01, pl01g)
    ca = kq["ca"] + ca_dac_ta(kq)
    bo = {
        "bo_ca_mau": "bo-ca-mau-twd-diem-cong-viec",
        "goi": "twd/diem-cong-viec",
        "gia_lap": True,
        "nguoi_ky": None,
        "yeu_cau_dat": 1.0,
        "trang_thai_ky": "NHAP",
        "mac_dinh": {"XEP_LOAI_CHUOI": {"co_quy_khong_hoan_thanh": False, "loai_tru": False,
                                        "rang_buoc_tap_the_thap_nhat": False, "rang_buoc_tap_the_ha_xuat_sac": False}},
        "ca": ca,
    }
    dau = ("# Bộ ca mẫu GIẢ LẬP của gói twd/diem-cong-viec (DL-16). TỆP SINH TỰ ĐỘNG bằng data/golden/twd/sinh_ca_twd.py\n"
           "# từ Phụ lục 01 và 01G của bộ tài liệu bàn giao UEB (dữ liệu giả định của tài liệu, không phải hồ sơ thật);\n"
           "# không sửa tay, sửa tập lệnh rồi sinh lại. Kỳ vọng lấy từ cột kết quả của bảng tính hoặc từ kết quả mong đợi\n"
           "# ghi trong Đặc tả (11.12 điểm h, 32.4, Phụ lục C), không tính bằng bộ chạy ca. Chưa được nghiệp vụ ký.\n"
           "# Chạy: python -I data/golden/chay_ca_mau.py data/policy-pack\n")
    DICH.write_text(dau + yaml.safe_dump(bo, allow_unicode=True, sort_keys=False, width=200, default_flow_style=None), encoding="utf-8")
    print(f"Đã ghi {DICH.relative_to(GOC)}: {len(ca)} ca")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
